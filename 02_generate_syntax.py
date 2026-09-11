"""
STAGE 2: Generate SPSS (.sps) syntax from a finalized mapping Excel file.

Takes the mapping file (after your team has reviewed/fixed any flagged rows
from Stage 1) plus the original .sav (for value labels, which live in the
.sav's metadata rather than the mapping sheet) and writes a ready-to-run
.sps syntax file: RENAME VARIABLES, VARIABLE LABELS, VALUE LABELS.

USAGE:
    python 02_generate_syntax.py <mapping.xlsx> <original.sav> <output.sps>
"""

import sys
import argparse
import openpyxl
import pyreadstat


def sps_quote(text):
    """Escape single quotes for SPSS string literals."""
    if text is None:
        text = ""
    return str(text).replace("'", "''")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("mapping_xlsx")
    parser.add_argument("original_sav")
    parser.add_argument("output_sps")
    args = parser.parse_args()

    # --- load finalized mapping ---
    wb = openpyxl.load_workbook(args.mapping_xlsx, data_only=True)
    ws = wb["Mapping"] if "Mapping" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    # columns: Raw Variable, Suggested Rename, Rename Needs Review, Rename Reason,
    #          Raw Variable Label, Cleaned Variable Label, Label Needs Review, Label Reason,
    #          Has Value Labels, Value Labels (preview)
    keep_rows = []
    for r in rows:
        raw_name, new_name = r[0], r[1]
        if not raw_name or new_name == "[EXCLUDED]":
            continue
        clean_label = r[5] if len(r) > 5 else ""
        keep_rows.append((raw_name, new_name, clean_label))

    # --- SAFETY NET (defense in depth): refuse to emit RENAME syntax that
    # would collide two different raw variables onto the same new name.
    # This is checked here too, independent of Stage 1, because this script
    # must never trust an upstream file blindly. ---
    from collections import defaultdict
    owners = defaultdict(list)
    for raw_name, new_name, _ in keep_rows:
        owners[new_name].append(raw_name)
    collisions = {k: v for k, v in owners.items() if len(v) > 1}
    if collisions:
        print("ERROR: mapping file contains name collisions — refusing to generate syntax.")
        print("The following new names are used by more than one raw variable:")
        for new_name, raws in collisions.items():
            print(f"  '{new_name}' <- {raws}")
        print("\nFix these rows in the mapping file (give each raw variable a unique new name) and re-run.")
        sys.exit(1)

    # --- pull value labels straight from the original .sav (source of truth) ---
    _, meta = pyreadstat.read_sav(args.original_sav)
    value_labels = meta.variable_value_labels

    lines = []
    lines.append("* Auto-generated SPSS syntax.")
    lines.append("* Generated from mapping file: {}".format(args.mapping_xlsx))
    lines.append("* Source data file: {}".format(args.original_sav))
    lines.append("")

    # --- RENAME VARIABLES ---
    renames = [(r[0], r[1]) for r in keep_rows if r[0] != r[1]]
    if renames:
        lines.append("* --- Rename variables ---.")
        lines.append("RENAME VARIABLES")
        for old, new in renames:
            lines.append(f"  ({old} = {new})")
        lines.append("  .")
        lines.append("")

    # --- VARIABLE LABELS ---
    lines.append("* --- Variable labels ---.")
    lines.append("VARIABLE LABELS")
    for raw_name, new_name, clean_label in keep_rows:
        if clean_label:
            lines.append(f"  {new_name} '{sps_quote(clean_label)}'")
    lines.append("  .")
    lines.append("")

    # --- VALUE LABELS ---
    lines.append("* --- Value labels ---.")
    name_map = {r[0]: r[1] for r in keep_rows}
    any_value_labels = False
    for raw_name, new_name, _ in keep_rows:
        vlabels = value_labels.get(raw_name)
        if not vlabels:
            continue
        any_value_labels = True
        lines.append(f"VALUE LABELS {new_name}")
        for code, label in sorted(vlabels.items(), key=lambda x: x[0]):
            code_str = str(int(code)) if float(code).is_integer() else str(code)
            lines.append(f"  {code_str} '{sps_quote(label)}'")
        lines.append("  .")
        lines.append("")

    if not any_value_labels:
        lines.append("* (no value-labeled variables found).")
        lines.append("")

    lines.append("EXECUTE.")

    with open(args.output_sps, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    print(f"Wrote {args.output_sps}")
    print(f"Variables renamed: {len(renames)}")
    print(f"Variables with labels written: {sum(1 for r in keep_rows if r[2])}")
    print(f"Variables with value labels: {sum(1 for r in keep_rows if value_labels.get(r[0]))}")


if __name__ == "__main__":
    main()
