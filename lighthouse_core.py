"""
lighthouse_core.py

Shared logic for the Lighthouse Studio -> SPSS mapping/syntax automation.
Both the CLI scripts (01_generate_mapping.py, 02_generate_syntax.py) and the
Streamlit app (app.py) import from this module, so there is exactly one
implementation of the rules to keep in sync.
"""

import re
from collections import defaultdict

import pyreadstat
import openpyxl
from openpyxl.styles import PatternFill, Font

# ---------------------------------------------------------------------------
# CONFIG (reusable across studies — edit/extend as new patterns appear)
# ---------------------------------------------------------------------------

DEFAULT_SYSTEM_FIELD_OVERRIDES = {
    "sys_RespNum":     ("Respondent_No", "Respondent Number"),
    "sys_StartTime":   ("Start_Time", "Start Time"),
    "sys_EndTime":     ("End_Time", "End Time"),
    "sys_ElapsedTime": ("Elapsed_Time", "Elapsed Time"),
}

DEFAULT_MANUAL_OVERRIDES = {
    "check_version": ("Version", "Version"),
    "lang":          ("Language", "Language"),
    "Coun":          ("Country", "Country"),
    "Password":      ("Password", "Password"),
}

PIPING_REPLACEMENTS = {
    r"\[%\s*ListLabel\(\s*Brand\s*,\s*1\s*\)\s*%\]": "Vivo V19",
    r"\[%\s*ListLabel\(\s*B4heading\s*,\s*\d+\s*\)\s*%\]": "",
}

COUNTRY_CODE_MAP = {
    "M": "Myanmar",
    "B": "Bangladesh",
}

EXCLUDE_PATTERNS = [
    r"^for[A-Z]",
    r"^sys_pagetime_",
]

COLUMN_TOPIC_MAP = {
    ("C6", "c1"): "PMI",
    ("C6", "c2"): "MHI",
}

# ---------------------------------------------------------------------------
# LABEL CLEANING
# ---------------------------------------------------------------------------

NOTE_PATTERNS = [
    r"\[Note:.*?\]",
    r"\[Interviewer note:.*?\]",
    r"Note:\s*\(Interviewer[^)]*\)",
    r"\(Interviewer[^)]*\)",
    r"Interviewer need to speak out the options",
]

TAG_PATTERNS = [
    r"\?\s*\((Single|Multiple)\s+Answer\)\s*$",
    r"\s*\((Single|Multiple)\s+Answer\)\s*$",
]

LEADING_CODE_PATTERN = re.compile(r"^\s*[A-Za-z0-9_]+\s*-\s*")


def resolve_piping(text):
    for pattern, replacement in PIPING_REPLACEMENTS.items():
        text = re.sub(pattern, replacement, text)
    return text


def clean_label(raw_label):
    """Returns (cleaned_text, needs_review_flag, review_reason)"""
    if not raw_label:
        return "", True, "empty raw label"

    text = raw_label
    text = LEADING_CODE_PATTERN.sub("", text)
    text = resolve_piping(text)

    unresolved_piping = re.findall(r"\[%.*?%\]", text)

    for pat in TAG_PATTERNS:
        text = re.sub(pat, "", text)

    review_reason = []
    for pat in NOTE_PATTERNS:
        if re.search(pat, text, flags=re.IGNORECASE):
            text = re.sub(pat, "", text, flags=re.IGNORECASE)

    text = re.sub(r"\s{2,}", " ", text).strip()

    needs_review = False
    if unresolved_piping:
        needs_review = True
        review_reason.append(f"unresolved piping: {unresolved_piping}")
    if "Interviewer" in text or "interviewer" in text:
        needs_review = True
        review_reason.append("possible leftover interviewer note")
    if "[" in text or "]" in text:
        needs_review = True
        review_reason.append("possible leftover bracket note")

    return text, needs_review, "; ".join(review_reason)


# ---------------------------------------------------------------------------
# VARIABLE RENAMING
# ---------------------------------------------------------------------------

GRID_RC_PATTERN = re.compile(r"^(?P<base>.+)_r(?P<row>\d+)_c(?P<col>\d+)$")
X_LOOP_PATTERN = re.compile(r"^(?P<base>.+?)x(?P<loop>\d+)(?P<suffix>_\d+.*)?$")
X_COUNTRY_PATTERN = re.compile(r"^(?P<base>.+?)x(?P<loop>\d+)(?P<country>[A-Z])(?P<suffix>_.*)?$")
COUNTRY_LETTER_SUFFIX_PATTERN = re.compile(r"^(?P<base>.+?)(?P<country>[A-Z])$")
COLUMN_CODE_PATTERN = re.compile(r"^(?P<base>[A-Za-z0-9]+?)(?P<country>[A-Z])_(?P<col>c\d+)$")


def is_excluded(name):
    return any(re.search(pat, name) for pat in EXCLUDE_PATTERNS)


def suggest_renames(raw_names):
    """Returns dict raw_name -> (suggested_new_name, needs_review, reason)"""
    results = {}

    rc_groups = defaultdict(list)
    for name in raw_names:
        m = GRID_RC_PATTERN.match(name)
        if m:
            rc_groups[m.group("base")].append(name)

    collapse_targets = defaultdict(list)
    for name in raw_names:
        m = X_LOOP_PATTERN.match(name)
        if m and m.group("suffix") and not X_COUNTRY_PATTERN.match(name):
            collapsed = f"{m.group('base')}{m.group('suffix')}"
            collapse_targets[collapsed].append(name)
    x_loop_collapse_is_safe = {k: len(v) for k, v in collapse_targets.items()}

    for name in raw_names:
        if is_excluded(name):
            results[name] = ("[EXCLUDED]", False, "matches exclude pattern — dropped from output")
            continue

        if name in DEFAULT_SYSTEM_FIELD_OVERRIDES:
            new_name, _ = DEFAULT_SYSTEM_FIELD_OVERRIDES[name]
            results[name] = (new_name, False, "system field dictionary")
            continue
        if name in DEFAULT_MANUAL_OVERRIDES:
            new_name, _ = DEFAULT_MANUAL_OVERRIDES[name]
            results[name] = (new_name, False, "manual override dictionary")
            continue

        m = GRID_RC_PATTERN.match(name)
        if m:
            base, row = m.group("base"), m.group("row")
            siblings = rc_groups[base]
            new_name = base if len(siblings) == 1 else f"{base}_{row}"
            results[name] = (new_name, False, "grid _r{n}_c pattern")
            continue

        m = X_COUNTRY_PATTERN.match(name)
        if m and m.group("country") in COUNTRY_CODE_MAP:
            base, loop, country = m.group("base"), m.group("loop"), m.group("country")
            suffix = m.group("suffix") or ""
            country_name = COUNTRY_CODE_MAP[country]
            new_name = f"{base}.{loop}{suffix}.{country_name}"
            results[name] = (new_name, True, "country-loop pattern — VERIFY dot-notation convention")
            continue

        m = X_LOOP_PATTERN.match(name)
        if m and m.group("suffix"):
            base, loop = m.group("base"), m.group("loop")
            collapsed = f"{base}{m.group('suffix')}"
            if x_loop_collapse_is_safe.get(collapsed) == 1:
                results[name] = (collapsed, True, "x-loop pattern — VERIFY convention matches this question type")
            else:
                safe_name = f"{base}_x{loop}{m.group('suffix')}"
                results[name] = (safe_name, True,
                                  "x-loop pattern COLLIDES across pages (same item # reused with different "
                                  "content) — kept loop number in name to avoid overwriting data. "
                                  "Confirm your team's real naming convention for this case.")
            continue

        m = COLUMN_CODE_PATTERN.match(name)
        if m and m.group("country") in COUNTRY_CODE_MAP:
            base, country, col = m.group("base"), m.group("country"), m.group("col")
            topic = COLUMN_TOPIC_MAP.get((base, col))
            if topic:
                new_name = f"{base}.{COUNTRY_CODE_MAP[country]}_{topic}"
                results[name] = (new_name, False, "column-topic dictionary")
            else:
                new_name = f"{base}.{COUNTRY_CODE_MAP[country]}_{col}"
                results[name] = (new_name, True, f"column code '{col}' not in topic dictionary — add mapping")
            continue

        m = COUNTRY_LETTER_SUFFIX_PATTERN.match(name)
        if m and m.group("country") in COUNTRY_CODE_MAP and len(name) > 2:
            base, country = m.group("base"), m.group("country")
            new_name = f"{base}.{COUNTRY_CODE_MAP[country]}"
            results[name] = (new_name, True, "trailing country-letter — VERIFY not a false positive")
            continue

        results[name] = (name, False, "no rename rule matched — kept as-is")

    # FINAL SAFETY NET: never let two different raw variables map to the same new name.
    new_name_owners = defaultdict(list)
    for raw, (new_name, _, _) in results.items():
        if new_name != "[EXCLUDED]":
            new_name_owners[new_name].append(raw)

    for new_name, owners in new_name_owners.items():
        if len(owners) > 1:
            for raw in owners:
                results[raw] = (
                    raw, True,
                    f"COLLISION BLOCKED: rule would have renamed {len(owners)} different variables "
                    f"({', '.join(owners[:5])}{'...' if len(owners) > 5 else ''}) to '{new_name}'. "
                    f"Reverted to raw name — resolve manually before running syntax."
                )

    return results


def load_codebook_appends(path_or_workbook):
    """
    Optional supplementary [Variable, Option Text] source. Accepts a file path
    or an already-opened openpyxl workbook (Streamlit passes file-like objects
    that openpyxl can load directly too).
    """
    if not path_or_workbook:
        return {}
    if isinstance(path_or_workbook, openpyxl.Workbook):
        wb = path_or_workbook
    else:
        wb = openpyxl.load_workbook(path_or_workbook, data_only=True)
    ws = wb.active
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            out[str(row[0]).strip()] = str(row[1]).strip() if row[1] else ""
    return out


# ---------------------------------------------------------------------------
# STAGE 1: build the mapping workbook
# ---------------------------------------------------------------------------

MAPPING_HEADERS = [
    "Raw Variable", "Suggested Rename", "Rename Needs Review", "Rename Reason",
    "Raw Variable Label", "Cleaned Variable Label", "Label Needs Review", "Label Reason",
    "Has Value Labels", "Value Labels (preview)"
]


def build_mapping_workbook(sav_path, codebook_path=None):
    """
    sav_path: path to a .sav file on disk.
    codebook_path: optional path (or file-like) to a [Variable, Option Text] xlsx.
    Returns (openpyxl.Workbook, stats_dict).
    """
    df, meta = pyreadstat.read_sav(sav_path)
    raw_names = list(meta.column_names)
    labels = meta.column_names_to_labels
    value_labels = meta.variable_value_labels
    codebook_appends = load_codebook_appends(codebook_path)

    renames = suggest_renames(raw_names)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mapping"
    ws.append(MAPPING_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    review_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    excluded_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

    n_review = 0
    n_excluded = 0
    for name in raw_names:
        new_name, rn_review, rn_reason = renames[name]

        if new_name == "[EXCLUDED]":
            n_excluded += 1
            ws.append([name, "[EXCLUDED]", "", rn_reason, "", "", "", "", "", ""])
            for col in range(1, len(MAPPING_HEADERS) + 1):
                ws.cell(row=ws.max_row, column=col).fill = excluded_fill
            continue

        raw_label = labels.get(name, "")
        clean, lbl_review, lbl_reason = clean_label(raw_label)

        if name in codebook_appends and codebook_appends[name]:
            clean = f"{clean} : {codebook_appends[name]}"
        elif lbl_reason == "" and re.search(r"_r\d+_c\d+$|_\d+$", name) and ":" not in clean:
            lbl_review = True
            lbl_reason = "grid item — per-option text not in .sav; provide a codebook or add manually"

        vlabels = value_labels.get(name)
        has_vlabels = "Yes" if vlabels else "No"
        vlabel_preview = "; ".join(f"{k}={v}" for k, v in list(vlabels.items())[:4]) if vlabels else ""
        if vlabels and len(vlabels) > 4:
            vlabel_preview += " ..."

        row_idx = ws.max_row + 1
        ws.append([
            name, new_name, "REVIEW" if rn_review else "", rn_reason,
            raw_label, clean, "REVIEW" if lbl_review else "", lbl_reason,
            has_vlabels, vlabel_preview
        ])
        if rn_review or lbl_review:
            n_review += 1
            for col in range(1, len(MAPPING_HEADERS) + 1):
                ws.cell(row=row_idx, column=col).fill = review_fill

    for i, w in enumerate([20, 24, 14, 30, 45, 45, 14, 30, 12, 40], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    kept = len(raw_names) - n_excluded
    stats = {
        "total": len(raw_names),
        "excluded": n_excluded,
        "kept": kept,
        "flagged": n_review,
        "auto_resolved": kept - n_review,
    }
    return wb, stats


# ---------------------------------------------------------------------------
# STAGE 2: build the .sps syntax text
# ---------------------------------------------------------------------------

def sps_quote(text):
    if text is None:
        text = ""
    return str(text).replace("'", "''")


def build_syntax_text(mapping_workbook, sav_path):
    """
    mapping_workbook: an openpyxl Workbook (already loaded) of the finalized mapping.
    sav_path: path to the ORIGINAL .sav (value labels are pulled from here).
    Returns (syntax_text, stats_dict). Raises ValueError on name collisions.
    """
    ws = mapping_workbook["Mapping"] if "Mapping" in mapping_workbook.sheetnames else mapping_workbook.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    keep_rows = []
    for r in rows:
        raw_name, new_name = r[0], r[1]
        if not raw_name or new_name == "[EXCLUDED]":
            continue
        clean_lbl = r[5] if len(r) > 5 else ""
        keep_rows.append((raw_name, new_name, clean_lbl))

    owners = defaultdict(list)
    for raw_name, new_name, _ in keep_rows:
        owners[new_name].append(raw_name)
    collisions = {k: v for k, v in owners.items() if len(v) > 1}
    if collisions:
        detail = "\n".join(f"  '{k}' <- {v}" for k, v in collisions.items())
        raise ValueError(
            "Mapping file contains name collisions — refusing to generate syntax.\n"
            f"{detail}\n"
            "Fix these rows (give each raw variable a unique new name) and try again."
        )

    _, meta = pyreadstat.read_sav(sav_path)
    value_labels = meta.variable_value_labels

    lines = []
    lines.append("* Auto-generated SPSS syntax.")
    lines.append("")

    renames = [(r[0], r[1]) for r in keep_rows if r[0] != r[1]]
    if renames:
        lines.append("* --- Rename variables ---.")
        lines.append("RENAME VARIABLES")
        for old, new in renames:
            lines.append(f"  ({old} = {new})")
        lines.append("  .")
        lines.append("")

    lines.append("* --- Variable labels ---.")
    lines.append("VARIABLE LABELS")
    for raw_name, new_name, clean_lbl in keep_rows:
        if clean_lbl:
            lines.append(f"  {new_name} '{sps_quote(clean_lbl)}'")
    lines.append("  .")
    lines.append("")

    lines.append("* --- Value labels ---.")
    any_value_labels = False
    for raw_name, new_name, _ in keep_rows:
        vlabels = value_labels.get(raw_name)
        if not vlabels:
            continue
        any_value_labels = True
        lines.append(f"VALUE LABELS {new_name}")
        for code, label in sorted(vlabels.items(), key=lambda x: x[0]):
            code_str = str(int(code)) if float(code).is_integer() else str(code)
            lines.append(f"  {code_str} '{sps_quote(label)}'")
        lines.append("  .")
        lines.append("")

    if not any_value_labels:
        lines.append("* (no value-labeled variables found).")
        lines.append("")

    lines.append("EXECUTE.")

    stats = {
        "renamed": len(renames),
        "labeled": sum(1 for r in keep_rows if r[2]),
        "value_labeled": sum(1 for r in keep_rows if value_labels.get(r[0])),
    }
    return "\n".join(lines), stats
