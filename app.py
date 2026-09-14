"""
Streamlit app for the SPSS mapping-file generator.

Run with:
    streamlit run app.py

Requires generate_mapping.py to be in the same folder — this app is just a
UI wrapper around its build_workbook() function; all the actual parsing/
renaming/labeling logic lives there and is unchanged.
"""
import io
import tempfile
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from generate_mapping import build_workbook

st.set_page_config(page_title="SPSS Mapping File Generator", layout="wide")

st.title("SPSS Mapping File Generator")
st.caption(
    "Upload a raw .sav data file and the Word questionnaire for the same study "
    "to get a draft Variable Label / Value Label mapping file."
)

col1, col2 = st.columns(2)
with col1:
    sav_file = st.file_uploader("Raw data file (.sav)", type=["sav"])
with col2:
    qnr_file = st.file_uploader("Questionnaire (.docx)", type=["docx"])

generate = st.button("Generate mapping file", type="primary", disabled=not (sav_file and qnr_file))

if generate:
    with tempfile.TemporaryDirectory() as tmpdir:
        sav_path = Path(tmpdir) / sav_file.name
        qnr_path = Path(tmpdir) / qnr_file.name
        sav_path.write_bytes(sav_file.getvalue())
        qnr_path.write_bytes(qnr_file.getvalue())

        with st.spinner("Parsing questionnaire and raw data, building mapping..."):
            try:
                wb = build_workbook(str(sav_path), str(qnr_path))
            except Exception as e:
                st.error(f"Failed to generate mapping: {e}")
                st.stop()

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        st.session_state["mapping_bytes"] = buf.getvalue()
        st.session_state["mapping_ready"] = True

if st.session_state.get("mapping_ready"):
    st.success("Mapping file generated.")

    st.download_button(
        "Download mapping .xlsx",
        data=st.session_state["mapping_bytes"],
        file_name="draft_mapping.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    wb_preview = load_workbook(io.BytesIO(st.session_state["mapping_bytes"]), read_only=True)

    ws = wb_preview["Variable Label"]
    rows = list(ws.iter_rows(values_only=True))
    df = pd.DataFrame(rows[2:], columns=rows[1])  # skip title row, use header row

    total = len(df)
    flagged = df["Notes"].apply(lambda x: bool(x)).sum()

    m1, m2, m3 = st.columns(3)
    m1.metric("Total variables", total)
    m2.metric("Flagged for review", flagged)
    m3.metric("Auto-resolved", total - flagged)

    tab1, tab2 = st.tabs(["Flagged rows (review these first)", "All variables"])

    with tab1:
        flagged_df = df[df["Notes"].astype(bool)]
        st.dataframe(flagged_df, use_container_width=True, height=500)

    with tab2:
        st.dataframe(df, use_container_width=True, height=500)

    with st.expander("Value Label sheet preview"):
        ws2 = wb_preview["Value Label"]
        rows2 = list(ws2.iter_rows(values_only=True))
        df2 = pd.DataFrame(rows2[2:], columns=rows2[1])
        st.dataframe(df2, use_container_width=True, height=400)
